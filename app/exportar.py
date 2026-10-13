"""Exportar el resumen de un documento o una conversación a PDF o Word.

Primero se arma el contenido como una lista simple de bloques y después se convierte al
formato pedido. Cada bloque es {"tipo": "titulo" | "h2" | "p" | "li" | "meta", "texto": ...};
el texto puede llevar **negrita**.
"""
import html
import io
import re
from datetime import datetime

import pymupdf
from docx import Document
from docx.shared import Pt, RGBColor

from app import db

FORMATOS = {
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}


# ----------------------------------------------------------- contenido a exportar
def _lineas(texto: str) -> list[str]:
    """Líneas de una lista ('- algo') sin la viñeta."""
    return [re.sub(r"^\s*[-•*]\s*", "", l).strip() for l in texto.splitlines() if l.strip()]


def bloques_resumen(usuario: str, nombre: str) -> tuple[str, list[dict]]:
    """Resumen + puntos clave + glosario del documento. Devuelve (nombre de archivo, bloques)."""
    resumen = db.obtener_campo(usuario, nombre, "resumen")
    puntos = db.obtener_campo(usuario, nombre, "puntos_clave")
    glosario = db.obtener_campo(usuario, nombre, "glosario")
    if not (resumen or puntos or glosario):
        raise ValueError("Este documento aún no tiene resumen, puntos clave ni glosario.")
    bloques = [
        {"tipo": "titulo", "texto": f"Resumen de {nombre}"},
        {"tipo": "meta", "texto": f"Generado el {datetime.now():%d/%m/%Y %H:%M}"},
    ]
    if resumen:
        bloques += [{"tipo": "h2", "texto": "Resumen"}, {"tipo": "p", "texto": resumen}]
    if puntos:
        bloques.append({"tipo": "h2", "texto": "Puntos clave"})
        bloques += [{"tipo": "li", "texto": l} for l in _lineas(puntos)]
    if glosario:
        bloques.append({"tipo": "h2", "texto": "Glosario"})
        bloques += [{"tipo": "li", "texto": l} for l in _lineas(glosario)]
    return f"resumen_{nombre.rsplit('.', 1)[0]}", bloques


def bloques_chat(usuario: str, conversacion_id: int) -> tuple[str, list[dict]]:
    """Toda la conversación, con las fuentes de cada respuesta."""
    conversaciones = {c["id"]: c for c in db.listar_conversaciones(usuario)}
    if conversacion_id not in conversaciones:
        raise ValueError("Conversación no encontrada.")
    titulo = conversaciones[conversacion_id]["titulo"]
    bloques = [
        {"tipo": "titulo", "texto": f"Conversación: {titulo}"},
        {"tipo": "meta", "texto": f"Usuario: {usuario} · Exportada el {datetime.now():%d/%m/%Y %H:%M}"},
    ]
    for m in db.obtener_mensajes(conversacion_id):
        bloques.append({"tipo": "h2", "texto": "Usuario" if m["rol"] == "user" else "Asistente"})
        bloques.append({"tipo": "p", "texto": m["texto"]})
        fuentes = m.get("fuentes") or []
        if fuentes:
            citas = "; ".join(f"[{f.get('n', i + 1)}] {f['documento']}, pág. {f['pagina']}"
                              for i, f in enumerate(fuentes))
            bloques.append({"tipo": "meta", "texto": f"Fuentes: {citas}"})
    return "conversacion", bloques


# ------------------------------------------------------------------------- PDF
_CSS = """
body { font-family: sans-serif; font-size: 10.5pt; }
h1 { font-size: 20pt; color: #1f3a68; }
h2 { font-size: 13pt; color: #1f3a68; margin-top: 14pt; }
p { margin-top: 3pt; margin-bottom: 3pt; }
li { margin-bottom: 3pt; }
.meta { color: #666666; font-size: 9pt; }
"""


def _a_html(texto: str) -> str:
    """Escapa el texto y convierte **negrita** en <b>."""
    seguro = html.escape(texto)
    seguro = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", seguro)
    seguro = re.sub(r"(?<!\*)\*(?!\*)(.+?)(?<!\*)\*(?!\*)", r"<i>\1</i>", seguro)
    return seguro.replace("\n", "<br>")


def a_pdf(bloques: list[dict]) -> bytes:
    """Convierte los bloques en un PDF (varias páginas si hace falta)."""
    partes, en_lista = [], False
    for b in bloques:
        if b["tipo"] != "li" and en_lista:
            partes.append("</ul>")
            en_lista = False
        t = _a_html(b["texto"])
        if b["tipo"] == "titulo":
            partes.append(f"<h1>{t}</h1>")
        elif b["tipo"] == "h2":
            partes.append(f"<h2>{t}</h2>")
        elif b["tipo"] == "meta":
            partes.append(f'<p class="meta">{t}</p>')
        elif b["tipo"] == "li":
            if not en_lista:
                partes.append("<ul>")
                en_lista = True
            partes.append(f"<li>{t}</li>")
        else:
            partes.append(f"<p>{t}</p>")
    if en_lista:
        partes.append("</ul>")

    historia = pymupdf.Story(html="".join(partes), user_css=_CSS)
    salida = io.BytesIO()
    escritor = pymupdf.DocumentWriter(salida)
    hoja = pymupdf.paper_rect("a4")
    area = hoja + (50, 50, -50, -50)  # márgenes
    hay_mas = True
    while hay_mas:
        dispositivo = escritor.begin_page(hoja)
        hay_mas, _ = historia.place(area)
        historia.draw(dispositivo)
        escritor.end_page()
    escritor.close()
    return salida.getvalue()


# ------------------------------------------------------------------------ Word
def _con_negrita(parrafo, texto: str):
    """Agrega el texto al párrafo respetando **negrita**."""
    for trozo in re.split(r"(\*\*.+?\*\*)", texto):
        if trozo.startswith("**") and trozo.endswith("**") and len(trozo) > 4:
            parrafo.add_run(trozo[2:-2]).bold = True
        elif trozo:
            parrafo.add_run(trozo)


def a_docx(bloques: list[dict]) -> bytes:
    """Convierte los bloques en un documento de Word."""
    doc = Document()
    for b in bloques:
        if b["tipo"] == "titulo":
            doc.add_heading(b["texto"], level=0)
        elif b["tipo"] == "h2":
            doc.add_heading(b["texto"], level=2)
        elif b["tipo"] == "meta":
            run = doc.add_paragraph().add_run(b["texto"])
            run.italic = True
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
        elif b["tipo"] == "li":
            _con_negrita(doc.add_paragraph(style="List Bullet"), b["texto"])
        else:
            _con_negrita(doc.add_paragraph(), b["texto"])
    salida = io.BytesIO()
    doc.save(salida)
    return salida.getvalue()


def exportar(bloques: list[dict], formato: str) -> bytes:
    if formato == "pdf":
        return a_pdf(bloques)
    if formato == "docx":
        return a_docx(bloques)
    raise ValueError("Formato no válido: usa 'pdf' o 'docx'.")
